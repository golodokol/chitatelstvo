# -*- coding: utf-8 -*-
"""Build Word TZ for letter-module voice actor sheet."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[1]
MD = ROOT / "docs" / "early-courses" / "09-voice-actor-letters-module1.md"
OUT = ROOT / "docs" / "early-courses" / "09-voice-actor-letters-module1.docx"


def _set_run_font(run, *, bold: bool = False, size: float = 11) -> None:
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Times New Roman"
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), "Times New Roman")


def _add_para(doc: Document, text: str, *, bold: bool = False, size: float = 11) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    _set_run_font(run, bold=bold, size=size)
    p.paragraph_format.space_after = Pt(6)


def _clean_cell(text: str) -> str:
    text = text.strip()
    text = re.sub(r"^`+|`+$", "", text)
    text = text.replace("**", "")
    return text


def _add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        _set_run_font(run, bold=True, size=10)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(_clean_cell(val))
            _set_run_font(run, size=10)
    doc.add_paragraph()


def main() -> None:
    lines = MD.read_text(encoding="utf-8").splitlines()
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(1.5)
        section.right_margin = Cm(1.5)

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        if line.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(line[2:].strip())
            _set_run_font(run, bold=True, size=16)
            i += 1
            continue
        if line.startswith("## "):
            _add_para(doc, line[3:].strip(), bold=True, size=14)
            i += 1
            continue
        if line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append(lines[i])
                i += 1
            parsed = []
            for raw in block:
                cells = [c.strip() for c in raw.strip().strip("|").split("|")]
                if all(re.fullmatch(r":?-{3,}:?", c or "") for c in cells):
                    continue
                parsed.append(cells)
            if parsed:
                headers = parsed[0]
                rows = parsed[1:]
                # Skip empty "how to speak" first column header row with blank first cell
                if headers == ["", ""] and rows:
                    # key-value table
                    kv = [[r[0], r[1]] for r in rows if len(r) >= 2]
                    _add_table(doc, ["Параметр", "Значение"], kv)
                else:
                    _add_table(doc, headers, rows)
            continue
        if line.startswith("- "):
            text = line[2:].replace("**", "").replace("`", "")
            p = doc.add_paragraph(style="List Bullet")
            p.clear()
            run = p.add_run(text)
            _set_run_font(run, size=11)
            i += 1
            continue
        text = line.replace("**", "").replace("`", "")
        _add_para(doc, text, size=11)
        i += 1

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
